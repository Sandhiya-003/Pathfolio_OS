"""Unit tests for ExtractionService against *real* file formats.

The earlier version of this test suite only ever uploaded .txt fixtures,
which meant a real bug (pypdf.PdfReader being used as a context manager,
and a python-docx import shadowing bug) shipped unnoticed -- every real PDF
and DOCX upload was silently failing. These tests generate actual .pdf and
.docx files on disk so extraction is exercised against the real formats
users will actually upload.
"""

from docx import Document as DocxDocument
from reportlab.pdfgen import canvas

from app.services.extraction_service import extraction_service


def _make_pdf(path, lines):
    c = canvas.Canvas(str(path))
    y = 750
    for line in lines:
        c.drawString(100, y, line)
        y -= 40
    c.save()


def _make_docx(path, paragraphs):
    doc = DocxDocument()
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.save(str(path))


def test_extract_real_pdf_returns_readable_text(tmp_path):
    pdf_path = tmp_path / "certificate.pdf"
    _make_pdf(pdf_path, [
        "Certificate of Completion",
        "Python for Data Science - Infosys Springboard",
        "Date: 15 March 2024",
    ])

    result = extraction_service.extract(str(pdf_path))

    assert result.get("error") is None
    assert result["file_type"] == "pdf"
    assert "Certificate of Completion" in result["text"]
    assert result["word_count"] > 0
    assert result["pages"] == 1


def test_extract_real_docx_returns_readable_text(tmp_path):
    docx_path = tmp_path / "report.docx"
    _make_docx(docx_path, [
        "Project Report",
        "FinVoice AI is a voice-first banking assistant.",
        "Technologies: FastAPI, React, Groq",
    ])

    result = extraction_service.extract(str(docx_path))

    assert result.get("error") is None
    assert result["file_type"] == "docx"
    assert "FinVoice AI" in result["text"]
    assert result["word_count"] > 0


def test_extract_docx_includes_table_content(tmp_path):
    docx_path = tmp_path / "with_table.docx"
    doc = DocxDocument()
    doc.add_paragraph("Skills Matrix")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Python"
    table.rows[0].cells[1].text = "Advanced"
    doc.save(str(docx_path))

    result = extraction_service.extract(str(docx_path))

    assert "Python" in result["text"]
    assert "Advanced" in result["text"]


def test_legacy_doc_gives_actionable_error_not_silent_failure(tmp_path):
    fake_doc_path = tmp_path / "old_resume.doc"
    fake_doc_path.write_bytes(b"not a real OLE2 .doc file")

    result = extraction_service.extract(str(fake_doc_path))

    assert result["text"] == ""
    assert "error" in result
    assert ".docx" in result["error"]  # tells the user what to do, not just "failed"


def test_extract_plain_text_file(tmp_path):
    txt_path = tmp_path / "notes.txt"
    txt_path.write_text("Internship completion letter for Summer 2024.")

    result = extraction_service.extract(str(txt_path))

    assert result["text"] == "Internship completion letter for Summer 2024."
    assert result.get("error") is None


def test_upload_a_real_pdf_end_to_end(client, auth_headers, tmp_path):
    """Full-stack regression test for the exact bug the user hit: uploading
    a real (non-.txt) PDF through the actual API, not just calling the
    service directly."""
    pdf_path = tmp_path / "aws_certificate.pdf"
    _make_pdf(pdf_path, [
        "AWS Cloud Practitioner Certificate",
        "Awarded to the bearer for completing AWS Cloud fundamentals.",
        "Date: June 2024",
    ])

    with open(pdf_path, "rb") as f:
        res = client.post(
            "/api/upload/single",
            headers=auth_headers,
            files={"file": ("aws_certificate.pdf", f, "application/pdf")},
        )

    assert res.status_code == 200
    body = res.json()
    assert body["success"] is True, body.get("message")
    assert body["category"] == "certification"


def test_upload_a_real_docx_end_to_end(client, auth_headers, tmp_path):
    """Full-stack regression test for the DOCX import-shadowing bug."""
    docx_path = tmp_path / "project_report.docx"
    _make_docx(docx_path, [
        "Smart EV Helper - Project Report",
        "A navigation tool for EV charging in connectivity dead zones.",
        "Built with Next.js, FastAPI, and Leaflet.js.",
    ])

    with open(docx_path, "rb") as f:
        res = client.post(
            "/api/upload/single",
            headers=auth_headers,
            files={"file": ("project_report.docx", f,
                             "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )

    assert res.status_code == 200
    body = res.json()
    assert body["success"] is True, body.get("message")